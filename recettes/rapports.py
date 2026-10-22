from datetime import date
from decimal import Decimal
from html import escape
from io import BytesIO
import re

from django.db.models import Sum
from django.db.models.functions import ExtractYear
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .alertes import calculer_ecarts, execution_budgetaire, recouvrement_par_secteur
from .formats import ar, pct
from .models import Anomalie, Categorie, ObjectifBudgetaire, ParametrageSysteme, Recette
from .prevision import avertissements_modele, prevu_contre_realise, previsions_par_categorie

MOIS_NOMS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
             "septembre", "octobre", "novembre", "décembre"]


def mois_depuis_texte(texte):
    try:
        annee, mois = (texte or "").split("-")
        return date(int(annee), int(mois), 1)
    except (ValueError, TypeError):
        return None


def _libelle_mois(d):
    return f"{MOIS_NOMS[d.month - 1].capitalize()} {d.year}"

def texte_sur(valeur):
    if isinstance(valeur, str) and valeur[:1] in ("=", "+", "-", "@", "\t", "\r"):
        if re.fullmatch(r"[+-]?[\d\s,.]+%?", valeur):
            return valeur
        return "'" + valeur
    return valeur

def comparaison_exercices(annees: int = 3, annee_reference: int | None = None) -> dict:
    lignes_brutes = (
        Recette.objects.annotate(annee=ExtractYear("mois"))
        .values("annee", "categorie_id", "categorie__nom", "categorie__chapitre")
        .annotate(total=Sum("montant"))
        .order_by("annee")
    )
    toutes_annees = sorted({
        ligne["annee"] for ligne in lignes_brutes
        if not annee_reference or ligne["annee"] <= annee_reference
    })
    if not toutes_annees:
        return {
            "annees": [], "lignes": [], "lignes_chapitres": [], "totaux_generaux": {},
            "variations_generales": {}, "avertissement_annee_partielle": None,
            "projection_totale": None, "tendance_totale": None, "mois_saisis": 0,
            "annees_comparees": [], "cellules_totaux_generaux": [], "cellules_variations_generales": [],
        }
    annees_retenues = toutes_annees[-annees:]

    par_categorie = {}
    for ligne in lignes_brutes:
        if ligne["annee"] not in annees_retenues:
            continue
        cle = ligne["categorie_id"]
        par_categorie.setdefault(cle, {"categorie": ligne["categorie__nom"], "chapitre": ligne["categorie__chapitre"], "totaux": {}})
        par_categorie[cle]["totaux"][ligne["annee"]] = ligne["total"]

    def _variations(totaux):
        variations = {}
        for precedente, courante in zip(annees_retenues, annees_retenues[1:]):
            avant = totaux.get(precedente)
            apres = totaux.get(courante)
            if avant and avant != 0 and apres is not None:
                variations[courante] = round(float((apres - avant) / avant * 100), 1)
            else:
                variations[courante] = None
        return variations

    lignes = []
    for donnees in par_categorie.values():
        totaux = {annee: donnees["totaux"].get(annee, Decimal("0")) for annee in annees_retenues}
        variations = _variations(totaux)
        lignes.append(
            {
                "categorie": donnees["categorie"],
                "chapitre": donnees["chapitre"],
                "totaux": totaux,
                "variations": variations,
                "cellules_totaux": [totaux[a] for a in annees_retenues],
                "cellules_variations": [variations.get(a) for a in annees_retenues[1:]],
            }
        )
    lignes.sort(key=lambda l: l["categorie"])

    totaux_generaux = {
        annee: sum((l["totaux"][annee] for l in lignes), Decimal("0")) for annee in annees_retenues
    }
    variations_generales = _variations(totaux_generaux)

    annee_courante = timezone.localdate().year
    avertissement = None
    mois_saisis = 12
    if annee_courante in annees_retenues:
        mois_saisis = (
            Recette.objects.filter(mois__year=annee_courante)
            .values_list("mois", flat=True)
            .distinct()
            .count()
        )
        if mois_saisis < 12:
            avertissement = annee_courante

    derniere = annees_retenues[-1]
    facteur = Decimal(12) / Decimal(mois_saisis) if avertissement and mois_saisis else Decimal(1)
    for l in lignes:
        l["projection"] = (l["totaux"][derniere] * facteur).quantize(Decimal("1"))
        precedent = l["totaux"].get(annees_retenues[-2]) if len(annees_retenues) > 1 else None
        l["tendance"] = (
            round(float((l["projection"] - precedent) / precedent * 100), 1) if precedent else None
        )
    par_chapitre = {}
    for l in lignes:
        c = par_chapitre.setdefault(l["chapitre"], {"chapitre": l["chapitre"], "totaux": {a: Decimal("0") for a in annees_retenues}})
        for a in annees_retenues:
            c["totaux"][a] += l["totaux"][a]
    lignes_chapitres = []
    for code in sorted(par_chapitre):
        c = par_chapitre[code]
        c["variations"] = _variations(c["totaux"])
        c["projection"] = (c["totaux"][derniere] * facteur).quantize(Decimal("1"))
        prec = c["totaux"].get(annees_retenues[-2]) if len(annees_retenues) > 1 else None
        c["tendance"] = round(float((c["projection"] - prec) / prec * 100), 1) if prec else None
        lignes_chapitres.append(c)
    projection_totale = (totaux_generaux[derniere] * facteur).quantize(Decimal("1"))
    precedent_total = totaux_generaux.get(annees_retenues[-2]) if len(annees_retenues) > 1 else None
    tendance_totale = (
        round(float((projection_totale - precedent_total) / precedent_total * 100), 1)
        if precedent_total else None
    )

    return {
        "annees": annees_retenues,
        "annees_comparees": annees_retenues[1:],
        "lignes": lignes,
        "lignes_chapitres": lignes_chapitres,
        "totaux_generaux": totaux_generaux,
        "variations_generales": variations_generales,
        "cellules_totaux_generaux": [totaux_generaux[a] for a in annees_retenues],
        "cellules_variations_generales": [variations_generales.get(a) for a in annees_retenues[1:]],
        "avertissement_annee_partielle": avertissement,
        "mois_saisis": mois_saisis,
        "projection_totale": projection_totale,
        "tendance_totale": tendance_totale,
        "nb_categories": Categorie.objects.count(),
    }

def _pied_de_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
    canvas.setLineWidth(0.5)
    canvas.line(doc.leftMargin, 1.25 * cm, A4[0] - doc.rightMargin, 1.25 * cm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(doc.leftMargin, 0.8 * cm, "PréviRecettes · Montants en Ariary (Ar)")
    canvas.drawRightString(A4[0] - doc.rightMargin, 0.8 * cm, f"Page {doc.page}")
    canvas.restoreState()

def _styles_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="PdfTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=18, leading=22, textColor=colors.HexColor("#17202A"),
        alignment=TA_LEFT, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="PdfMeta", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8.5, leading=11, textColor=colors.HexColor("#64748B"),
        spaceAfter=14,
    ))
    styles.add(ParagraphStyle(
        name="PdfHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=11, leading=14, textColor=colors.HexColor("#047857"),
        spaceBefore=7, spaceAfter=7,
    ))
    styles.add(ParagraphStyle(
        name="PdfNote", parent=styles["Italic"], fontName="Helvetica-Oblique",
        fontSize=8, leading=10, textColor=colors.HexColor("#64748B"),
    ))
    return styles


def _ajouter_tableau(elements, donnees, largeurs, compact=False):
    elements.append(_tableau(donnees, largeurs, compact=compact))
    elements.append(Spacer(1, 0.8 * cm))


def _dernier_mois_de(annee):
    qs = Recette.objects.all()
    if annee:
        qs = qs.filter(mois__year=annee)
    return qs.order_by("-mois").values_list("mois", flat=True).first()


def _section_liasse(elements, styles, annee, mois):
    mois = mois or _dernier_mois_de(annee)
    if mois is None:
        elements.append(Paragraph("Aucune recette enregistrée pour cette période.", styles["PdfNote"]))
        return
    elements.append(Paragraph(f"Période : {_libelle_mois(mois)}", styles["PdfHeading"]))
    recettes = Recette.objects.filter(mois=mois).select_related(
        "categorie", "secteur"
    ).order_by("categorie__chapitre", "categorie__code", "secteur__nom")
    objectifs = {
        ligne["categorie_id"]: ligne["t"]
        for ligne in ObjectifBudgetaire.objects.filter(
            mois=mois, statut=ObjectifBudgetaire.Statut.INSCRIT, secteur__isnull=True
        ).values("categorie_id").annotate(t=Sum("montant_cible"))
    }
    realise = {}
    for recette in recettes:
        entree = realise.setdefault(
            recette.categorie_id, {"categorie": recette.categorie, "total": Decimal("0")}
        )
        entree["total"] += recette.montant

    elements.append(Paragraph("Exécution par compte", styles["PdfHeading"]))
    donnees = [["Chapitre", "Compte", "Réalisé (Ar)", "Objectif voté (Ar)", "Écart"]]
    total_realise = total_objectif = Decimal("0")
    for entree in sorted(realise.values(), key=lambda e: (e["categorie"].chapitre, e["categorie"].code)):
        cible = objectifs.get(entree["categorie"].pk)
        ecart = float((entree["total"] - cible) / cible * 100) if cible else None
        donnees.append([
            entree["categorie"].chapitre, entree["categorie"].libelle_complet, ar(entree["total"]),
            ar(cible) if cible else "—", pct(ecart, 1, signe=True) if ecart is not None else "—",
        ])
        total_realise += entree["total"]
        total_objectif += cible or Decimal("0")
    ecart_total = float((total_realise - total_objectif) / total_objectif * 100) if total_objectif else None
    donnees.append([
        "", "TOTAL", ar(total_realise), ar(total_objectif) if total_objectif else "—",
        pct(ecart_total, 1, signe=True) if ecart_total is not None else "—",
    ])
    _ajouter_tableau(elements, donnees, [1.6 * cm, 6.6 * cm, 3.2 * cm, 3.6 * cm, 2.0 * cm])

    elements.append(Paragraph("Détail des encaissements du mois", styles["PdfHeading"]))
    donnees = [["Date", "Bordereau", "Compte", "Secteur", "Montant (Ar)"]]
    nombre_recettes = recettes.count()
    for recette in recettes.order_by("date_imputation", "id")[:60]:
        donnees.append([
            f"{recette.date_imputation:%d/%m/%Y}" if recette.date_imputation else _libelle_mois(recette.mois),
            recette.reference, recette.categorie.nom, recette.secteur.nom, ar(recette.montant),
        ])
    if nombre_recettes > 60:
        donnees.append(["", "", f"... {nombre_recettes - 60} autres lignes (voir l'export Excel)", "", ""])
    _ajouter_tableau(elements, donnees, [2.2 * cm, 2.6 * cm, 5.4 * cm, 3.4 * cm, 3.4 * cm], compact=True)

    elements.append(Paragraph("Anomalies ouvertes", styles["PdfHeading"]))
    donnees = [["Réf.", "Secteur", "Motif", "Montant (Ar)", "Échéance"]]
    for anomalie in Anomalie.objects.select_related("secteur").filter(brouillon=False, traitee=False):
        donnees.append([
            anomalie.reference, anomalie.secteur.nom, anomalie.get_motif_display(), ar(anomalie.montant),
            f"{anomalie.date_echeance:%d/%m/%Y}" if anomalie.date_echeance else "—",
        ])
    if len(donnees) == 1:
        donnees.append(["—", "Aucune anomalie ouverte", "—", "—", "—"])
    _ajouter_tableau(elements, donnees, [2.4 * cm, 3.6 * cm, 4.6 * cm, 3.4 * cm, 3.0 * cm])


def _section_synthese(elements, styles, annee, parametres):
    elements.append(Paragraph(f"Trajectoire mensuelle de l'exercice {annee}", styles["PdfHeading"]))
    realises = {
        ligne["mois"]: ligne["t"]
        for ligne in Recette.objects.filter(mois__year=annee).values("mois").annotate(t=Sum("montant"))
    }
    objectifs = {
        ligne["mois"]: ligne["t"]
        for ligne in ObjectifBudgetaire.objects.filter(
            mois__year=annee, statut=ObjectifBudgetaire.Statut.INSCRIT, secteur__isnull=True
        ).values("mois").annotate(t=Sum("montant_cible"))
    }
    donnees = [["Mois", "Réalisé (Ar)", "Objectif voté (Ar)", "Cumul réalisé (Ar)"]]
    cumul = Decimal("0")
    for numero in range(1, 13):
        mois = date(annee, numero, 1)
        if mois not in realises and mois not in objectifs:
            continue
        cumul += realises.get(mois, Decimal("0"))
        donnees.append([
            _libelle_mois(mois), ar(realises.get(mois)) if mois in realises else "—",
            ar(objectifs.get(mois)) if mois in objectifs else "—", ar(cumul),
        ])
    if len(donnees) == 1:
        donnees.append(["—", "Aucune donnée", "—", "—"])
    _ajouter_tableau(elements, donnees, [4.0 * cm, 4.2 * cm, 4.6 * cm, 4.2 * cm])

    execution = execution_budgetaire(annee)
    elements.append(Paragraph(
        f"Taux d'exécution budgétaire : {pct(execution['taux_pct'], 1) if execution['taux_pct'] is not None else '—'} "
        f"(réalisé {ar(execution['total_realise'])} pour un objectif de {ar(execution['total_cible'])}).",
        styles["PdfNote"],
    ))
    elements.append(Spacer(1, 0.6 * cm))
    elements.append(Paragraph(f"Comparaison entre exercices (jusqu'à {annee})", styles["PdfHeading"]))
    comparaison = comparaison_exercices(3, annee)
    if comparaison["annees"]:
        entete = ["Catégorie"] + [str(a) + " (Ar)" for a in comparaison["annees"]]
        entete += [f"Var. {a}" for a in comparaison["annees"][1:]]
        donnees = [entete]
        for ligne in comparaison["lignes"]:
            rangee = [ligne["categorie"]] + [ar(ligne["totaux"][a]) for a in comparaison["annees"]]
            rangee += [
                "—" if ligne["variations"].get(a) is None else pct(ligne["variations"][a], 1, signe=True)
                for a in comparaison["annees"][1:]
            ]
            donnees.append(rangee)
        rangee = ["TOTAL"] + [ar(comparaison["totaux_generaux"][a]) for a in comparaison["annees"]]
        rangee += [
            "—" if comparaison["variations_generales"].get(a) is None
            else pct(comparaison["variations_generales"][a], 1, signe=True)
            for a in comparaison["annees"][1:]
        ]
        donnees.append(rangee)
        _ajouter_tableau(elements, donnees, [5.2 * cm] + [2.35 * cm] * (len(donnees[0]) - 1))
        if comparaison["avertissement_annee_partielle"]:
            elements.append(Paragraph(
                f"Note : l'exercice {comparaison['avertissement_annee_partielle']} est encore incomplet ; "
                "les variations le concernant sont à lire avec prudence.", styles["PdfNote"],
            ))
            elements.append(Spacer(1, 0.4 * cm))
    else:
        elements.append(Paragraph("Aucune donnée exploitable.", styles["PdfNote"]))

    elements.append(Paragraph(f"Recouvrement par secteur (exercice {annee})", styles["PdfHeading"]))
    suivi = recouvrement_par_secteur(annee=annee)
    donnees = [["Secteur", "Total encaissé (Ar)", "Part", "Anomalies ouvertes"]]
    for ligne in suivi["lignes"]:
        donnees.append([
            ligne["secteur"].nom, ar(ligne["total"]), pct(ligne["part_pct"], 1),
            str(ligne["anomalies_ouvertes"]),
        ])
    if len(donnees) == 1:
        donnees.append(["—", "Aucune donnée", "—", "—"])
    _ajouter_tableau(elements, donnees, [4.7 * cm, 4.2 * cm, 2.1 * cm, 4 * cm])

    elements.append(Paragraph("Écarts entre objectifs votés et réalisé", styles["PdfHeading"]))
    donnees = [["Mois", "Catégorie", "Secteur", "Objectif voté (Ar)", "Réalisé (Ar)", "Écart"]]
    for ecart in calculer_ecarts(annee)[:20]:
        donnees.append([
            ecart.mois.strftime("%m/%Y"), ecart.categorie, ecart.secteur or "Consolidé",
            ar(ecart.cible), ar(ecart.realise), pct(ecart.ecart_pct, 1, signe=True),
        ])
    if len(donnees) == 1:
        donnees.append(["—", "Aucun objectif défini", "—", "—", "—", "—"])
    _ajouter_tableau(elements, donnees, [1.8 * cm, 4.4 * cm, 2.6 * cm, 3.0 * cm, 3.0 * cm, 2.2 * cm])

    elements.append(Paragraph(
        f"Prévisions statistiques à {parametres.mois_a_prevoir} mois "
        "(tendance linéaire, saisonnalité, facteurs économiques)", styles["PdfHeading"],
    ))
    elements.append(Paragraph(
        "Ces montants sont calculés à partir de l'historique, avec un intervalle de confiance à 95 %. "
        "Ils ne doivent pas être confondus avec les objectifs budgétaires votés.", styles["PdfNote"],
    ))
    elements.append(Spacer(1, 0.3 * cm))
    previsions = previsions_par_categorie()
    for message in avertissements_modele(previsions):
        elements.append(Paragraph(escape(message), styles["PdfNote"]))
    elements.append(Spacer(1, 0.3 * cm))
    donnees = [["Catégorie", "Mois", "Prévu (Ar)", "Borne basse (Ar)", "Borne haute (Ar)"]]
    for serie in previsions.values():
        for mois, valeur, bas, haut in zip(
            serie["mois_prevus"], serie["montants_prevus"], serie["bornes_basses"], serie["bornes_hautes"]
        ):
            donnees.append([serie["label"], mois, ar(valeur), ar(bas), ar(haut)])
    if len(donnees) == 1:
        donnees.append(["—", "Historique insuffisant", "—", "—", "—"])
    _ajouter_tableau(elements, donnees, [5.2 * cm, 2.0 * cm, 3.2 * cm, 3.2 * cm, 3.4 * cm], compact=True)

    historise = prevu_contre_realise()
    if historise:
        elements.append(Paragraph("Prévu contre réalisé (prévisions historisées)", styles["PdfHeading"]))
        donnees = [["Mois", "Catégorie", "Prévu (Ar)", "Réalisé (Ar)", "Écart"]]
        for ligne in historise[:30]:
            donnees.append([
                ligne["mois"].strftime("%m/%Y"), ligne["categorie"], ar(ligne["prevu"]),
                ar(ligne["realise"]),
                pct(ligne["ecart_pct"], 1, signe=True) if ligne["ecart_pct"] is not None else "—",
            ])
        _ajouter_tableau(elements, donnees, [2.0 * cm, 6.0 * cm, 3.2 * cm, 3.2 * cm, 2.6 * cm], compact=True)


def generer_rapport_pdf(type_doc: str = "LIASSE", annee: int | None = None, mois=None) -> BytesIO:
    parametres = ParametrageSysteme.charger()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2.2 * cm, bottomMargin=1.8 * cm,
        title="Rapport des recettes", author="PréviRecettes",
    )
    styles = _styles_pdf()
    annee = annee or (mois.year if mois else timezone.localdate().year)
    elements = []
    titres = {
        "LIASSE": "Liasse mensuelle — Exécution des recettes",
        "SYNTHESE": "Synthèse analytique des recettes",
        "ANOMALIES": "État des anomalies de recouvrement",
    }
    elements.append(Paragraph(titres.get(type_doc, titres["LIASSE"]), styles["PdfTitle"]))
    elements.append(Paragraph(
        f"Projet pédagogique, données fictives · Généré le {timezone.localdate():%d/%m/%Y} · Exercice {annee}",
        styles["PdfMeta"],
    ))
    if type_doc == "ANOMALIES":
        _section_anomalies(elements, styles)
    elif type_doc == "SYNTHESE":
        _section_synthese(elements, styles, annee, parametres)
    else:
        _section_liasse(elements, styles, annee, mois)
    doc.build(elements, onFirstPage=_pied_de_page, onLaterPages=_pied_de_page)
    buffer.seek(0)
    return buffer

def _section_anomalies(elements, styles):
    elements.append(Paragraph("Anomalies de recouvrement", styles["PdfHeading"]))
    donnees = [["Réf.", "Secteur", "Motif", "Montant (Ar)", "Échéance", "Statut"]]
    for a in Anomalie.objects.select_related("secteur", "categorie").filter(brouillon=False):
        donnees.append([
            a.reference, a.secteur.nom, a.get_motif_display(), ar(a.montant),
            f"{a.date_echeance:%d/%m/%Y}" if a.date_echeance else "—", a.statut,
        ])
    if len(donnees) == 1:
        donnees.append(["—", "Aucune anomalie", "—", "—", "—", "—"])
    table = _tableau(donnees, [2.2 * cm, 3.5 * cm, 4.1 * cm, 2.7 * cm, 2.3 * cm, 2.4 * cm])
    table.setStyle(_style_tableau())
    elements.append(table)


def _styler_feuille(ws):
    from openpyxl.styles import Font, PatternFill

    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="25272C")
    for colonne in ws.columns:
        ws.column_dimensions[colonne[0].column_letter].width = max(
            12, min(45, max(len(str(cellule.value or "")) for cellule in colonne) + 2)
        )


def _feuilles_synthese(wb, ws, annee):
    ws.title = "Trajectoire"
    ws.append(["Mois", "Réalisé (Ar)", "Objectif voté (Ar)", "Cumul réalisé (Ar)"])
    realises = {
        ligne["mois"]: ligne["t"]
        for ligne in Recette.objects.filter(mois__year=annee).values("mois").annotate(t=Sum("montant"))
    }
    objectifs = {
        ligne["mois"]: ligne["t"]
        for ligne in ObjectifBudgetaire.objects.filter(
            mois__year=annee, statut=ObjectifBudgetaire.Statut.INSCRIT, secteur__isnull=True
        ).values("mois").annotate(t=Sum("montant_cible"))
    }
    cumul = Decimal("0")
    for numero in range(1, 13):
        mois = date(annee, numero, 1)
        if mois in realises or mois in objectifs:
            cumul += realises.get(mois, Decimal("0"))
            ws.append([
                mois.strftime("%Y-%m"), float(realises.get(mois, 0)),
                float(objectifs.get(mois, 0)), float(cumul),
            ])

    feuille = wb.create_sheet("Secteurs")
    feuille.append(["Secteur", "Total encaissé (Ar)", "Part (%)", "Anomalies ouvertes"])
    for ligne in recouvrement_par_secteur(annee=annee)["lignes"]:
        feuille.append([
            ligne["secteur"].nom, float(ligne["total"]), ligne["part_pct"], ligne["anomalies_ouvertes"],
        ])

    feuille = wb.create_sheet("Écarts")
    feuille.append(["Mois", "Catégorie", "Secteur", "Objectif voté (Ar)", "Réalisé (Ar)", "Écart (%)"])
    for ecart in calculer_ecarts(annee):
        feuille.append([
            ecart.mois.strftime("%Y-%m"), ecart.categorie, ecart.secteur or "Consolidé",
            float(ecart.cible), float(ecart.realise), ecart.ecart_pct,
        ])

    feuille = wb.create_sheet("Prévisions")
    feuille.append(["Catégorie", "Mois", "Prévu (Ar)", "Borne basse (Ar)", "Borne haute (Ar)"])
    for serie in previsions_par_categorie().values():
        for mois, valeur, bas, haut in zip(
            serie["mois_prevus"], serie["montants_prevus"], serie["bornes_basses"], serie["bornes_hautes"]
        ):
            feuille.append([serie["label"], mois, valeur, bas, haut])


def generer_rapport_xlsx(type_doc: str = "LIASSE", annee: int | None = None, mois=None) -> BytesIO:
    from openpyxl import Workbook

    annee = annee or (mois.year if mois else timezone.localdate().year)
    wb = Workbook()
    ws = wb.active
    if type_doc == "ANOMALIES":
        ws.title = "Anomalies"
        ws.append(["Référence", "Secteur", "Compte", "Motif", "Montant (Ar)", "Date de constat", "Échéance", "Statut"])
        for anomalie in Anomalie.objects.select_related("secteur", "categorie").filter(brouillon=False):
            ws.append([
                anomalie.reference, anomalie.secteur.nom, anomalie.categorie.libelle_complet,
                anomalie.get_motif_display(), float(anomalie.montant), anomalie.date_constat,
                anomalie.date_echeance, anomalie.statut,
            ])
    elif type_doc == "SYNTHESE":
        _feuilles_synthese(wb, ws, annee)
    else:
        mois = mois or _dernier_mois_de(annee)
        ws.title = "Liasse " + (mois.strftime("%Y-%m") if mois else "vide")
        ws.append(["Référence", "Date", "Chapitre", "Compte", "Secteur", "Mode", "Montant (Ar)", "Libellé"])
        qs = Recette.objects.select_related("categorie", "secteur").order_by("date_imputation", "id")
        for recette in qs.filter(mois=mois) if mois else qs.none():
            ws.append([
                recette.reference, recette.date_imputation, recette.categorie.chapitre,
                recette.categorie.libelle_complet, recette.secteur.nom,
                recette.get_mode_reglement_display(), float(recette.montant), texte_sur(recette.libelle),
            ])
    for feuille in wb.worksheets:
        _styler_feuille(feuille)
    tampon = BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return tampon

def _tableau(donnees, largeurs, compact=False):
    taille = 6.2 if compact else 7.4
    interligne = 7.4 if compact else 9
    cellules = []
    for index_ligne, ligne in enumerate(donnees):
        style = ParagraphStyle(
            name=f"PdfCell{index_ligne}{compact}",
            fontName="Helvetica-Bold" if index_ligne == 0 else "Helvetica",
            fontSize=taille,
            leading=interligne,
            textColor=colors.white if index_ligne == 0 else colors.HexColor("#25313C"),
            alignment=TA_CENTER if index_ligne == 0 else TA_LEFT,
        )
        cellules.append([
            cell if isinstance(cell, Paragraph) else Paragraph(escape(str(cell)).replace("\n", "<br/>"), style)
            for cell in ligne
        ])
    tableau = Table(cellules, colWidths=largeurs, repeatRows=1, hAlign="LEFT")
    tableau.setStyle(_style_tableau())
    return tableau

def _style_tableau():
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#25272C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#10B981")),
            ("GRID", (0, 1), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
    )
