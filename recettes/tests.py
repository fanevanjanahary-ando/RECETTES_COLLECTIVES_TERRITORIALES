import smtplib
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.contrib import admin
from django.contrib.messages import get_messages
from django.core import mail
from django.core.files.base import ContentFile
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db.utils import IntegrityError
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .alertes import alertes_significatives, calculer_ecarts, recouvrement_par_secteur
from .models import (
    Anomalie,
    Categorie,
    Notification,
    ObjectifBudgetaire,
    ParametrageSysteme,
    Profil,
    RapportGenere,
    Recette,
    Secteur,
    TentativeConnexion,
)
from .forms import RecetteForm
from .notifications import (
    _creer_et_envoyer,
    notifier_anomalies_en_attente,
    notifier_ecarts_significatifs,
    notifier_echeances_recouvrement,
)
from .prevision import _t_critique_95, previsions_par_categorie
from .rapports import comparaison_exercices, generer_rapport_xlsx, texte_sur

class BaseDonnees(TestCase):
    def setUp(self):
        self.taxes = Categorie.objects.create(nom="Taxes", ouverte_aux_agents=True)
        self.subventions = Categorie.objects.create(nom="Subventions", ouverte_aux_agents=False)
        self.nord = Secteur.objects.create(nom="Secteur Nord", code="NORD")
        self.sud = Secteur.objects.create(nom="Secteur Sud", code="SUD")

        self.agent = User.objects.create_user("agent", "agent@example.test", "MotDePasse!2026")
        self.agent.profil.role = Profil.Role.AGENT
        self.agent.profil.secteur = self.nord
        self.agent.profil.save()

        self.responsable = User.objects.create_user("respo", "respo@example.test", "MotDePasse!2026")
        self.responsable.profil.role = Profil.Role.RESPONSABLE
        self.responsable.profil.save()

        self.admin = User.objects.create_user("admin", "admin@example.test", "MotDePasse!2026")
        self.admin.profil.role = Profil.Role.ADMIN
        self.admin.profil.save()

class Point1_RestrictionCategorieAgent(BaseDonnees):

    def test_agent_ne_peut_pas_poster_une_categorie_interdite(self):
        self.client.force_login(self.agent)
        reponse = self.client.post(
            reverse("ajouter_recette"),
            {
                "categorie": self.subventions.pk,
                "secteur": self.nord.pk,
                "date_imputation": "2026-01-05",
                "montant": "1000.00",
                "mode_reglement": "VIREMENT",
                "libelle": "Test",
            },
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(Recette.objects.filter(categorie=self.subventions).exists())

    def test_agent_peut_saisir_une_categorie_autorisee(self):
        self.client.force_login(self.agent)
        reponse = self.client.post(
            reverse("ajouter_recette"),
            {
                "categorie": self.taxes.pk,
                "secteur": self.nord.pk,
                "date_imputation": "2026-01-05",
                "montant": "1000.00",
                "mode_reglement": "VIREMENT",
                "libelle": "Test",
            },
        )
        self.assertEqual(reponse.status_code, 302)
        self.assertTrue(Recette.objects.filter(categorie=self.taxes).exists())

    def test_responsable_peut_saisir_toutes_les_categories(self):
        self.client.force_login(self.responsable)
        reponse = self.client.post(
            reverse("ajouter_recette"),
            {
                "categorie": self.subventions.pk,
                "secteur": self.nord.pk,
                "date_imputation": "2026-01-05",
                "montant": "500.00",
                "mode_reglement": "VIREMENT",
                "libelle": "Test",
            },
        )
        self.assertEqual(reponse.status_code, 302)

class Point9_MontantNegatif(BaseDonnees):

    def test_full_clean_refuse_un_montant_negatif(self):
        recette = Recette(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("-500"))
        with self.assertRaises(ValidationError):
            recette.full_clean()

    def test_la_base_refuse_un_montant_negatif_meme_sans_validation(self):
        with self.assertRaises(IntegrityError):
            Recette.objects.create(
                categorie=self.taxes, secteur=self.nord, mois=date(2026, 2, 1), montant=Decimal("-500")
            )

    def test_objectif_negatif_refuse(self):
        objectif = ObjectifBudgetaire(categorie=self.taxes, mois=date(2026, 1, 1), montant_cible=Decimal("-10"))
        with self.assertRaises(ValidationError):
            objectif.full_clean()

class Point10_ComptesDemoAvecEmail(TestCase):

    def test_les_comptes_de_demo_ont_une_adresse(self):
        from django.core.management import call_command

        call_command("generer_donnees_fictives", "--mois", "4", "--force", verbosity=0)
        for username in ("admin_demo", "responsable_demo", "elu_demo", "agent_demo"):
            utilisateur = User.objects.get(username=username)
            self.assertTrue(utilisateur.email, f"{username} n'a pas d'adresse e-mail")
            self.assertIn("@", utilisateur.email)

class Point3_NotificationsReelles(BaseDonnees):

    def test_un_ecart_significatif_declenche_un_email_et_une_notification(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("1000"))
        ObjectifBudgetaire.objects.create(categorie=self.taxes, mois=date(2026, 1, 1), montant_cible=Decimal("2000"))

        envoyees = notifier_ecarts_significatifs()
        self.assertGreater(envoyees, 0)
        self.assertGreater(len(mail.outbox), 0)
        self.assertTrue(Notification.objects.filter(type=Notification.Type.ECART).exists())

    def test_pas_de_doublon_si_on_relance(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("1000"))
        ObjectifBudgetaire.objects.create(categorie=self.taxes, mois=date(2026, 1, 1), montant_cible=Decimal("2000"))
        premier = notifier_ecarts_significatifs()
        second = notifier_ecarts_significatifs()
        self.assertGreater(premier, 0)
        self.assertEqual(second, 0)

class Point2_SuiviParSecteur(BaseDonnees):

    def test_le_recouvrement_est_ventile_par_secteur(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("3000"))
        Recette.objects.create(categorie=self.taxes, secteur=self.sud, mois=date(2026, 1, 1), montant=Decimal("1000"))
        suivi = recouvrement_par_secteur()
        parts = {l["secteur"].nom: l["part_pct"] for l in suivi["lignes"]}
        self.assertEqual(parts["Secteur Nord"], 75.0)
        self.assertEqual(parts["Secteur Sud"], 25.0)

    def test_barres_de_repartition_utilisent_un_pourcentage_css_valide(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("3000"))
        Recette.objects.create(categorie=self.taxes, secteur=self.sud, mois=date(2026, 1, 1), montant=Decimal("1000"))
        self.client.force_login(self.responsable)

        reponse = self.client.get(reverse("suivi_secteurs"))

        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, 'style="width: 75.0%"')
        self.assertContains(reponse, 'style="width: 25.0%"')

    def test_anomalie_rattachee_a_un_secteur(self):
        anomalie = Anomalie.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1),
            description="Retard", signale_par=self.agent,
        )
        self.assertEqual(anomalie.secteur, self.nord)

class Point4_ModeleConfigurable(BaseDonnees):

    def test_la_page_de_parametrage_est_reservee_a_l_administrateur(self):
        self.client.force_login(self.responsable)
        reponse = self.client.get(reverse("parametrage"))
        self.assertEqual(reponse.status_code, 302)

    def test_l_administrateur_modifie_l_horizon_et_le_seuil(self):
        self.client.force_login(self.admin)
        reponse = self.client.post(
            reverse("parametrage"),
            {
                "mois_a_prevoir": 12,
                "historique_minimum": 4,
                "saisonnalite_minimum_mois": 24,
                "seuil_alerte_pct": "5.00",
                "jour_echeance_recouvrement": 15,
                "notifications_email_actives": "on",
            },
        )
        self.assertEqual(reponse.status_code, 302)
        parametres = ParametrageSysteme.charger()
        self.assertEqual(parametres.mois_a_prevoir, 12)
        self.assertEqual(parametres.seuil_alerte_pct, Decimal("5.00"))

    def test_le_seuil_configure_change_le_declenchement_des_alertes(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("1100"))
        ObjectifBudgetaire.objects.create(categorie=self.taxes, mois=date(2026, 1, 1), montant_cible=Decimal("1000"))

        parametres = ParametrageSysteme.charger()
        parametres.seuil_alerte_pct = Decimal("15.00")
        parametres.save()
        self.assertFalse(calculer_ecarts()[0].significatif)

        parametres.seuil_alerte_pct = Decimal("5.00")
        parametres.save()
        self.assertTrue(calculer_ecarts()[0].significatif)

class Point5_ComparaisonExercices(BaseDonnees):

    def test_variation_entre_deux_exercices(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2024, 1, 1), montant=Decimal("1000"))
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2025, 1, 1), montant=Decimal("1200"))
        resultat = comparaison_exercices()
        self.assertEqual(resultat["annees"], [2024, 2025])
        self.assertEqual(resultat["lignes"][0]["variations"][2025], 20.0)

class Point6_GestionDesComptes(BaseDonnees):

    def test_l_administrateur_voit_la_liste_des_comptes(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("liste_utilisateurs")).status_code, 200)

    def test_un_responsable_ne_peut_pas_gerer_les_comptes(self):
        self.client.force_login(self.responsable)
        self.assertEqual(self.client.get(reverse("liste_utilisateurs")).status_code, 302)

    def test_modification_du_role_et_de_l_email(self):
        self.client.force_login(self.admin)
        reponse = self.client.post(
            reverse("modifier_utilisateur", args=[self.agent.pk]),
            {
                "first_name": "Jean", "last_name": "Dupont",
                "email": "jean.dupont@example.test", "is_active": "on",
                "role": Profil.Role.RESPONSABLE, "recevoir_notifications": "on",
            },
        )
        self.assertEqual(reponse.status_code, 302)
        self.agent.refresh_from_db()
        self.assertEqual(self.agent.email, "jean.dupont@example.test")
        self.assertEqual(self.agent.profil.role, Profil.Role.RESPONSABLE)

    def test_desactivation_d_un_compte(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("basculer_utilisateur", args=[self.agent.pk]))
        self.agent.refresh_from_db()
        self.assertFalse(self.agent.is_active)

class Point7_ExportTableur(BaseDonnees):

    def test_export_recettes_csv(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("1000"))
        self.client.force_login(self.responsable)
        reponse = self.client.get(reverse("export_csv", args=["recettes"]))
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("text/csv", reponse["Content-Type"])
        contenu = reponse.content.decode("utf-8-sig")
        self.assertIn("Secteur", contenu)
        self.assertIn("Secteur Nord", contenu)

    def test_tous_les_exports_repondent(self):
        self.client.force_login(self.responsable)
        for jeu in ("recettes", "objectifs", "ecarts", "previsions", "exercices", "secteurs", "categories"):
            with self.subTest(jeu=jeu):
                self.assertEqual(self.client.get(reverse("export_csv", args=[jeu])).status_code, 200)

    def test_export_inconnu_renvoie_404(self):
        self.client.force_login(self.responsable)
        self.assertEqual(self.client.get(reverse("export_csv", args=["nimporte_quoi"])).status_code, 404)

    def test_page_rapports_accessible_au_responsable(self):
        self.client.force_login(self.responsable)
        reponse = self.client.get(reverse("rapports"))
        self.assertEqual(reponse.status_code, 200)
        self.assertIn(("categories", "Catégories (CSV)"), reponse.context["exports"])

class Point8_DistinctionPrevuObjectif(BaseDonnees):

    def test_le_vocabulaire_est_explicite_sur_le_tableau_de_bord(self):
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), montant=Decimal("1100"))
        ObjectifBudgetaire.objects.create(categorie=self.taxes, mois=date(2026, 1, 1), montant_cible=Decimal("2000"))
        self.client.force_login(self.responsable)
        contenu = self.client.get(reverse("dashboard")).content.decode()
        self.assertIn("projection", contenu)
        self.assertIn("Objectif voté", contenu)

    def test_la_page_budget_confronte_les_deux_notions(self):
        self.client.force_login(self.responsable)
        contenu = self.client.get(reverse("liste_objectifs")).content.decode()
        self.assertIn("Objectif voté confronté à la prévision du modèle", contenu)

class InscriptionEmailObligatoire(TestCase):

    def test_inscription_sans_email_refusee(self):
        reponse = self.client.post(
            reverse("inscription"),
            {
                "username": "nouveau", "email": "",
                "password1": "MotDePasse!2026", "password2": "MotDePasse!2026",
                "role": Profil.Role.RESPONSABLE,
            },
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(User.objects.filter(username="nouveau").exists())

class MoteurDePrevision(TestCase):

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Redevances")
        self.secteur = Secteur.objects.create(nom="Secteur unique")

    def _ajouter_serie(self, valeurs, depart=date(2025, 1, 1)):
        mois = depart
        for valeur in valeurs:
            Recette.objects.create(categorie=self.categorie, secteur=self.secteur, mois=mois, montant=Decimal(str(valeur)))
            mois = date(mois.year + 1, 1, 1) if mois.month == 12 else date(mois.year, mois.month + 1, 1)

    def test_extrapole_correctement_une_tendance_lineaire(self):
        self._ajouter_serie([1000, 1100, 1200, 1300, 1400, 1500])
        resultat = previsions_par_categorie(mois_a_prevoir=3, historique_minimum=3)
        serie = resultat[str(self.categorie.id)]
        self.assertEqual(serie["montants_prevus"], [1600.0, 1700.0, 1800.0])
        self.assertEqual(serie["mois_prevus"], ["2025-07", "2025-08", "2025-09"])

    def test_categorie_sous_le_seuil_est_ignoree(self):
        self._ajouter_serie([1000, 1100])
        resultat = previsions_par_categorie(mois_a_prevoir=3, historique_minimum=3)
        self.assertNotIn(str(self.categorie.id), resultat)

    def test_ne_prevoit_jamais_de_montant_negatif(self):
        self._ajouter_serie([500, 300, 100, 0, 0, 0])
        resultat = previsions_par_categorie(mois_a_prevoir=6, historique_minimum=3)
        serie = resultat[str(self.categorie.id)]
        self.assertTrue(all(v >= 0 for v in serie["montants_prevus"]))

class ConnexionParEmail(BaseDonnees):

    def test_connexion_avec_l_adresse_email(self):
        reponse = self.client.post(
            reverse("connexion"), {"username": "respo@example.test", "password": "MotDePasse!2026"}
        )
        self.assertEqual(reponse.status_code, 302)

    def test_connexion_avec_identifiant_acceptee(self):
        reponse = self.client.post(reverse("connexion"), {"username": "respo", "password": "MotDePasse!2026"})
        self.assertEqual(reponse.status_code, 302)


class DemandeHabilitation(TestCase):

    def test_inscription_cree_un_compte_inactif_en_attente(self):
        secteur = Secteur.objects.create(nom="Secteur Test", code="TST")
        reponse = self.client.post(
            reverse("inscription"),
            {
                "first_name": "Jean",
                "last_name": "Rakoto",
                "email": "jean.rakoto@collectivite.mg",
                "matricule": "RH-001",
                "role": "AGENT",
                "secteur": secteur.pk,
                "password1": "UnMotDePasse123",
                "password2": "UnMotDePasse123",
                "deontologie": "on",
            },
        )
        self.assertEqual(reponse.status_code, 200)
        utilisateur = User.objects.get(email="jean.rakoto@collectivite.mg")
        self.assertFalse(utilisateur.is_active)
        self.assertTrue(utilisateur.profil.en_attente)

    def test_administrateur_valide_l_habilitation(self):
        admin = User.objects.create_user("admin2", "admin2@example.test", "MotDePasse!2026")
        admin.profil.role = Profil.Role.ADMIN
        admin.profil.save()
        candidat = User.objects.create_user("candidat", "candidat@example.test", "x", is_active=False)
        candidat.profil.en_attente = True
        candidat.profil.save()

        self.client.force_login(admin)
        reponse = self.client.post(reverse("arbitrer_utilisateur", args=[candidat.pk]), {"decision": "valider"})
        self.assertEqual(reponse.status_code, 302)
        candidat.refresh_from_db()
        self.assertTrue(candidat.is_active)
        self.assertFalse(candidat.profil.en_attente)


class NomenclatureCategories(TestCase):

    def test_creation_d_un_compte_avec_code_et_chapitre(self):
        categorie = Categorie.objects.create(nom="Taxe de séjour", code="7362", chapitre="73")
        self.assertEqual(categorie.libelle_complet, "7362 · Taxe de séjour")

    def test_devise_ariary_dans_les_gabarits(self):
        admin = User.objects.create_user("admin3", "admin3@example.test", "MotDePasse!2026")
        admin.profil.role = Profil.Role.ADMIN
        admin.profil.save()
        self.client.force_login(admin)
        contenu = self.client.get(reverse("liste_categories")).content.decode()
        self.assertIn("Ar", contenu)
        self.assertNotIn("€", contenu)


class AnomaliesDeRecouvrement(BaseDonnees):

    def test_signalement_notifie_les_responsables(self):
        self.client.force_login(self.agent)
        reponse = self.client.post(
            reverse("signaler_anomalie"),
            {
                "categorie": self.taxes.pk,
                "secteur": self.nord.pk,
                "motif": "RETARD",
                "montant": "5000",
                "date_constat": "2026-01-05",
                "date_echeance": "2026-01-20",
                "description": "Retard de versement constaté en régie.",
            },
        )
        self.assertEqual(reponse.status_code, 302)
        anomalie = Anomalie.objects.get(secteur=self.nord)
        self.assertFalse(anomalie.brouillon)
        self.assertTrue(Notification.objects.filter(type=Notification.Type.ANOMALIE).exists())

    def test_traitement_reserve_a_la_gestion(self):
        anomalie = Anomalie.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1),
            montant=Decimal("100"), description="Test", signale_par=self.agent,
        )
        self.client.force_login(self.agent)
        reponse = self.client.post(reverse("traiter_anomalie", args=[anomalie.pk]))
        self.assertEqual(reponse.status_code, 302)
        anomalie.refresh_from_db()
        self.assertFalse(anomalie.traitee)

        self.client.force_login(self.responsable)
        reponse = self.client.post(reverse("traiter_anomalie", args=[anomalie.pk]))
        anomalie.refresh_from_db()
        self.assertTrue(anomalie.traitee)


class Cloisonnement(BaseDonnees):

    def setUp(self):
        super().setUp()
        media_temporaire = tempfile.TemporaryDirectory()
        self.addCleanup(media_temporaire.cleanup)
        parametrage_media = override_settings(MEDIA_ROOT=media_temporaire.name)
        parametrage_media.enable()
        self.addCleanup(parametrage_media.disable)
        self.agent.profil.secteur = self.nord
        self.agent.profil.save()
        self.recette_nord = Recette.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1),
            date_imputation=date(2026, 1, 5), montant=Decimal("1000"),
        )
        self.recette_sud = Recette.objects.create(
            categorie=self.taxes, secteur=self.sud, mois=date(2026, 1, 1),
            date_imputation=date(2026, 1, 6), montant=Decimal("2000"),
        )
        for recette in (self.recette_nord, self.recette_sud):
            recette.piece_jointe.save("bordereau.pdf", ContentFile(b"%PDF-1.4"), save=True)

    def test_la_liste_est_limitee_au_secteur_de_l_agent(self):
        self.client.force_login(self.agent)
        reponse = self.client.get(reverse("liste_recettes"))
        self.assertContains(reponse, self.recette_nord.reference)
        self.assertNotContains(reponse, self.recette_sud.reference)

    def test_le_detail_et_le_bordereau_d_un_autre_secteur_sont_refuses(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("detail_recette", args=[self.recette_sud.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("bordereau", args=[self.recette_sud.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("bordereau", args=[self.recette_nord.pk])).status_code, 200)

    def test_piece_jointe_absente_renvoie_404(self):
        recette = self.recette_nord
        recette.piece_jointe.storage.delete(recette.piece_jointe.name)
        self.client.force_login(self.responsable)

        reponse = self.client.get(reverse("bordereau", args=[recette.pk]))

        self.assertEqual(reponse.status_code, 404)

    def test_piece_d_anomalie_absente_renvoie_404(self):
        anomalie = Anomalie.objects.create(
            categorie=self.taxes,
            secteur=self.nord,
            mois=date(2026, 1, 1),
            description="Pièce à restaurer",
            signale_par=self.agent,
        )
        anomalie.piece_jointe.save("anomalie.pdf", ContentFile(b"%PDF-1.4"), save=True)
        anomalie.piece_jointe.storage.delete(anomalie.piece_jointe.name)
        self.client.force_login(self.agent)

        reponse = self.client.get(reverse("piece_anomalie", args=[anomalie.pk]))

        self.assertEqual(reponse.status_code, 404)

    def test_la_saisie_pour_un_autre_secteur_est_refusee(self):
        self.client.force_login(self.agent)
        reponse = self.client.post(
            reverse("ajouter_recette"),
            {
                "categorie": self.taxes.pk, "secteur": self.sud.pk, "date_imputation": "2026-02-05",
                "montant": "1000.00", "mode_reglement": "VIREMENT", "libelle": "Intrus",
            },
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(Recette.objects.filter(libelle="Intrus").exists())

    def test_l_export_de_l_agent_est_limite_a_son_secteur(self):
        self.client.force_login(self.agent)
        reponse = self.client.get(reverse("export_csv", args=["recettes"]))
        contenu = reponse.content.decode("utf-8-sig")
        self.assertIn(self.recette_nord.reference, contenu)
        self.assertNotIn(self.recette_sud.reference, contenu)

    def test_un_agent_sans_secteur_ne_voit_aucune_donnee_ni_secteur(self):
        self.agent.profil.secteur = None
        self.agent.profil.save()
        self.client.force_login(self.agent)

        reponse = self.client.get(reverse("liste_recettes"))
        self.assertEqual(reponse.context["total_selection"], Decimal("0"))
        self.assertEqual(reponse.context["secteurs"].count(), 0)

        export = self.client.get(reverse("export_csv", args=["recettes"]))
        contenu = export.content.decode("utf-8-sig")
        self.assertNotIn(self.recette_nord.reference, contenu)
        self.assertNotIn(self.recette_sud.reference, contenu)

        referentiels = self.client.get(reverse("api_referentiels")).json()
        self.assertEqual(referentiels["secteurs"], [])

    def test_le_dashboard_agent_est_limite_a_son_secteur(self):
        Recette.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 2, 1), montant=Decimal("1000")
        )
        Recette.objects.create(
            categorie=self.taxes, secteur=self.sud, mois=date(2026, 2, 1), montant=Decimal("9000")
        )
        ObjectifBudgetaire.objects.create(
            categorie=self.taxes, secteur=None, mois=date(2026, 2, 1), montant_cible=Decimal("50000")
        )
        ObjectifBudgetaire.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 2, 1), montant_cible=Decimal("1200")
        )
        Anomalie.objects.create(
            categorie=self.taxes, secteur=self.sud, mois=date(2026, 2, 1), montant=Decimal("500"),
            description="Dossier du secteur Sud.",
        )
        self.client.force_login(self.agent)

        reponse = self.client.get(reverse("dashboard"))
        self.assertEqual(reponse.context["total"], Decimal("2000"))
        self.assertEqual(reponse.context["execution"]["total_cible"], Decimal("1200"))
        self.assertEqual(reponse.context["anomalies_ouvertes_total"], 0)
        self.assertEqual(
            [ligne["secteur"].pk for ligne in reponse.context["suivi_secteurs"]["lignes"]],
            [self.nord.pk],
        )

    def test_le_suivi_global_est_reserve_aux_roles_budgetaires(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("suivi_secteurs")).status_code, 302)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertNotContains(dashboard, "Secteurs &amp; Recouvrement")

        self.client.force_login(self.responsable)
        self.assertEqual(self.client.get(reverse("suivi_secteurs")).status_code, 200)

    def test_l_elu_n_ouvre_pas_les_bordereaux(self):
        elu = User.objects.create_user("elu", "elu@example.test", "MotDePasse!2026")
        elu.profil.role = Profil.Role.ELU
        elu.profil.save()
        self.client.force_login(elu)
        self.assertNotEqual(self.client.get(reverse("bordereau", args=[self.recette_nord.pk])).status_code, 200)

    def test_parametres_d_url_non_numeriques_ne_provoquent_pas_d_erreur(self):
        self.client.force_login(self.responsable)
        for url in (
            "/recettes/?secteur=abc", "/recettes/?categorie=abc",
            "/anomalies/?secteur=abc", "/anomalies/?dossier=abc",
        ):
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_un_brouillon_d_autrui_reste_invisible(self):
        brouillon = Anomalie.objects.create(
            secteur=self.nord, categorie=self.taxes, mois=date(2026, 1, 1), montant=Decimal("500"),
            brouillon=True, signale_par=self.agent, description="Secret",
        )
        self.client.force_login(self.responsable)
        reponse = self.client.get(f"/anomalies/?dossier={brouillon.pk}")
        self.assertNotContains(reponse, "Secret")


class CorrectionsAudit(BaseDonnees):

    def test_le_rappel_d_echeance_charge_les_categories_actives(self):
        nombre = notifier_echeances_recouvrement(date(2026, 3, 15))
        self.assertGreater(nombre, 0)
        self.assertTrue(Notification.objects.filter(type=Notification.Type.ECHEANCE).exists())

    def test_la_promotion_en_superutilisateur_met_a_jour_le_profil(self):
        utilisateur = User.objects.create_user("super", "super@example.test", "MotDePasse!2026")
        utilisateur.is_superuser = True
        utilisateur.save(update_fields=["is_superuser"])
        utilisateur.profil.refresh_from_db()
        self.assertEqual(utilisateur.profil.role, Profil.Role.ADMIN)

    def test_echec_envoi_invitation_est_signale_sans_annuler_le_compte(self):
        self.client.force_login(self.admin)
        with patch("recettes.views.PasswordResetForm.save", side_effect=OSError("SMTP indisponible")):
            reponse = self.client.post(
                reverse("inviter_utilisateur"),
                {
                    "first_name": "Nouvel",
                    "last_name": "Agent",
                    "email": "nouvel.agent@example.test",
                    "role": Profil.Role.AGENT,
                },
                follow=True,
            )

        self.assertTrue(User.objects.filter(email="nouvel.agent@example.test").exists())
        messages_text = [str(message) for message in get_messages(reponse.wsgi_request)]
        self.assertTrue(any("mais l'envoi du lien a échoué" in message for message in messages_text))

    def test_un_email_de_notification_en_echec_peut_etre_reessaye(self):
        with patch("recettes.notifications.send_mail", side_effect=smtplib.SMTPException("SMTP indisponible")):
            with self.assertLogs("recettes.notifications", level="ERROR"):
                self.assertTrue(
                    _creer_et_envoyer(
                        self.admin, Notification.Type.SYSTEME, "Test", "Message", "TEST-RETRY"
                    )
                )

        notification = Notification.objects.get(cle_unicite="TEST-RETRY")
        self.assertFalse(notification.envoyee_par_email)

        with patch("recettes.notifications.send_mail", return_value=1) as send_mail:
            self.assertTrue(
                _creer_et_envoyer(
                    self.admin, Notification.Type.SYSTEME, "Test", "Message", "TEST-RETRY"
                )
            )

        notification.refresh_from_db()
        self.assertTrue(notification.envoyee_par_email)
        self.assertEqual(Notification.objects.filter(cle_unicite="TEST-RETRY").count(), 1)
        send_mail.assert_called_once()


class GardeFousComptes(BaseDonnees):

    def test_refuser_un_compte_actif_n_est_pas_possible(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("arbitrer_utilisateur", args=[self.agent.pk]), {"decision": "refuser"})
        self.assertTrue(User.objects.filter(pk=self.agent.pk).exists())

    def test_l_admin_ne_peut_pas_se_retrograder(self):
        self.client.force_login(self.admin)
        self.client.post(
            reverse("modifier_utilisateur", args=[self.admin.pk]),
            {
                "first_name": "A", "last_name": "B", "email": "admin@example.test",
                "role": Profil.Role.AGENT, "is_active": "on",
            },
        )
        self.admin.profil.refresh_from_db()
        self.assertEqual(self.admin.profil.role, Profil.Role.ADMIN)

    def test_le_dernier_admin_ne_peut_pas_etre_desactive(self):
        autre = User.objects.create_user("admin2", "admin2@example.test", "MotDePasse!2026")
        autre.profil.role = Profil.Role.ADMIN
        autre.profil.save()
        self.client.force_login(autre)
        self.client.post(reverse("basculer_utilisateur", args=[self.admin.pk]))
        self.client.force_login(self.admin)
        self.client.post(reverse("basculer_utilisateur", args=[autre.pk]))
        self.assertTrue(User.objects.get(pk=autre.pk).is_active)


class ConnexionLimitee(BaseDonnees):

    def test_apres_plusieurs_echecs_la_connexion_est_bloquee(self):
        for _ in range(5):
            self.client.post(reverse("connexion"), {"username": "respo", "password": "faux"})
        self.assertEqual(TentativeConnexion.objects.count(), 1)
        self.assertEqual(TentativeConnexion.objects.get().nombre, 5)
        reponse = self.client.post(
            reverse("connexion"), {"username": "respo", "password": "MotDePasse!2026"}
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertNotIn("_auth_user_id", self.client.session)


class DoubleComptageObjectifs(BaseDonnees):

    def test_un_objectif_de_secteur_ne_double_pas_le_realise(self):
        from .alertes import execution_budgetaire

        mois = date(2026, 1, 1)
        Recette.objects.create(categorie=self.taxes, secteur=self.nord, mois=mois, montant=Decimal("1000"))
        ObjectifBudgetaire.objects.create(categorie=self.taxes, mois=mois, montant_cible=Decimal("2000"))
        avant = execution_budgetaire(2026)
        ObjectifBudgetaire.objects.create(
            categorie=self.taxes, mois=mois, secteur=self.nord, montant_cible=Decimal("900")
        )
        apres = execution_budgetaire(2026)
        self.assertEqual(avant["total_cible"], apres["total_cible"])
        self.assertEqual(avant["total_realise"], apres["total_realise"])

    def test_plusieurs_types_de_budget_sont_possibles_pour_un_meme_mois(self):
        mois = date(2026, 1, 1)
        ObjectifBudgetaire.objects.create(
            categorie=self.taxes, mois=mois, montant_cible=Decimal("2000"),
            type_budget=ObjectifBudgetaire.TypeBudget.BP,
        )
        ObjectifBudgetaire.objects.create(
            categorie=self.taxes, mois=mois, montant_cible=Decimal("500"),
            type_budget=ObjectifBudgetaire.TypeBudget.DM,
        )


class PrevisionCalendrier(BaseDonnees):

    def test_les_mois_manquants_decalent_la_tendance(self):
        from .prevision import _modele

        mois = [
            date(2025, 1, 1), date(2025, 2, 1), date(2025, 3, 1),
            date(2025, 6, 1), date(2025, 7, 1),
        ]
        _, montants, bas, haut, saisonnalite = _modele(mois, [100, 110, 120, 130, 140], 2)
        self.assertEqual(montants, [145.07, 151.04])
        self.assertFalse(saisonnalite)
        self.assertTrue(all(b <= m <= h for b, m, h in zip(bas, montants, haut)))

    def test_un_facteur_economique_modifie_la_prevision(self):
        from .models import FacteurEconomique

        for i in range(6):
            Recette.objects.create(
                categorie=self.taxes, secteur=self.nord, mois=date(2025, 1 + i, 1),
                montant=Decimal(1000 + 10 * i),
            )
        base = previsions_par_categorie(
            mois_a_prevoir=1, historique_minimum=3
        )[str(self.taxes.pk)]["montants_prevus"][0]
        FacteurEconomique.objects.create(
            libelle="Hausse", mois_debut=date(2025, 7, 1), mois_fin=date(2025, 7, 1),
            effet_pct=Decimal("10"),
        )
        apres = previsions_par_categorie(
            mois_a_prevoir=1, historique_minimum=3
        )[str(self.taxes.pk)]["montants_prevus"][0]
        self.assertAlmostEqual(apres, round(base * 1.1, 2), places=1)


class PlafondEtFichiers(BaseDonnees):

    def test_le_plafond_d_encaissement_est_controle(self):
        self.taxes.type_perception = Categorie.Perception.DIRECTE
        self.taxes.save()
        self.nord.plafond_encaissement = Decimal("1000")
        self.nord.save()
        self.client.force_login(self.responsable)
        reponse = self.client.post(
            reverse("ajouter_recette"),
            {
                "categorie": self.taxes.pk, "secteur": self.nord.pk, "date_imputation": "2026-02-05",
                "montant": "5000.00", "mode_reglement": "ESPECES", "libelle": "Trop",
            },
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertFalse(Recette.objects.filter(libelle="Trop").exists())

    def test_un_faux_pdf_est_refuse(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .forms import valider_piece

        with self.assertRaises(ValidationError):
            valider_piece(SimpleUploadedFile("faux.pdf", b"<html><script>x</script></html>"))
        valider_piece(SimpleUploadedFile("vrai.pdf", b"%PDF-1.4 contenu"))


class ModifierRecetteFormulaire(BaseDonnees):

    def test_le_montant_est_prerempli_avec_un_point_et_les_boutons_existent(self):
        recette = Recette.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1),
            date_imputation=date(2026, 1, 5), montant=Decimal("3457250.00"),
        )
        self.client.force_login(self.responsable)
        reponse = self.client.get(reverse("modifier_recette", args=[recette.pk]))
        self.assertContains(reponse, 'value="3457250.00"')
        detail = self.client.get(reverse("detail_recette", args=[recette.pk]))
        self.assertContains(detail, reverse("supprimer_recette", args=[recette.pk]))


class CorrectionsAuditRegressions(BaseDonnees):

    def setUp(self):
        super().setUp()
        media_temporaire = tempfile.TemporaryDirectory()
        self.addCleanup(media_temporaire.cleanup)
        parametrage_media = override_settings(MEDIA_ROOT=media_temporaire.name)
        parametrage_media.enable()
        self.addCleanup(parametrage_media.disable)

    def test_date_imputation_future_refusee_par_le_formulaire_et_le_modele(self):
        demain = timezone.localdate() + timedelta(days=1)
        donnees = {
            "categorie": self.taxes.pk,
            "secteur": self.nord.pk,
            "date_imputation": demain.isoformat(),
            "montant": "100",
            "mode_reglement": "VIREMENT",
            "libelle": "Date future",
        }
        formulaire = RecetteForm(data=donnees, utilisateur=self.responsable)
        self.assertFalse(formulaire.is_valid())
        self.assertIn("date_imputation", formulaire.errors)

        recette = Recette(
            categorie=self.taxes, secteur=self.nord, date_imputation=demain, montant=Decimal("100")
        )
        with self.assertRaises(ValidationError):
            recette.full_clean()

    @patch("recettes.alertes.timezone.localdate", return_value=date(2026, 10, 22))
    def test_alerte_ecart_exclut_le_mois_en_cours(self, _localdate):
        for mois in (date(2026, 9, 1), date(2026, 10, 1)):
            ObjectifBudgetaire.objects.create(
                categorie=self.taxes, secteur=self.nord, mois=mois, montant_cible=Decimal("1000")
            )
            Recette.objects.create(
                categorie=self.taxes, secteur=self.nord, mois=mois, montant=Decimal("100")
            )

        alertes = alertes_significatives(2026)

        self.assertEqual([alerte.mois for alerte in alertes], [date(2026, 9, 1)])

    def test_api_agent_ne_recoit_que_les_categories_autorisees(self):
        self.client.force_login(self.agent)

        reponse = self.client.get(reverse("api_referentiels"))

        self.assertEqual(reponse.status_code, 200)
        noms = {categorie["nom"] for categorie in reponse.json()["categories"]}
        self.assertIn(self.taxes.nom, noms)
        self.assertNotIn(self.subventions.nom, noms)
        self.assertEqual(
            self.client.get(reverse("api_ecarts")).status_code,
            403,
        )

    def test_un_agent_ne_peut_pas_recevoir_un_compteur_global_d_anomalies(self):
        Anomalie.objects.create(
            categorie=self.taxes, secteur=self.nord, mois=date(2026, 1, 1), description="Nord"
        )
        Anomalie.objects.create(
            categorie=self.taxes, secteur=self.sud, mois=date(2026, 1, 1), description="Sud"
        )
        self.client.force_login(self.agent)

        reponse = self.client.get(reverse("signaler_anomalie"))

        self.assertEqual(reponse.context["nb_en_cours"], 1)

    def test_arbitrage_exige_une_decision_connue_et_conserve_un_refus(self):
        candidat = User.objects.create_user(
            "candidat-refuse", "candidat@example.test", "MotDePasse!2026", is_active=False
        )
        candidat.profil.en_attente = True
        candidat.profil.save(update_fields=["en_attente"])
        self.client.force_login(self.admin)

        invalide = self.client.post(
            reverse("arbitrer_utilisateur", args=[candidat.pk]), {"decision": "autre"}
        )
        candidat.refresh_from_db()
        self.assertEqual(invalide.status_code, 302)
        self.assertTrue(candidat.profil.en_attente)
        self.assertFalse(candidat.is_active)

        refuse = self.client.post(
            reverse("arbitrer_utilisateur", args=[candidat.pk]), {"decision": "refuser"}
        )
        candidat.refresh_from_db()
        self.assertEqual(refuse.status_code, 302)
        self.assertTrue(User.objects.filter(pk=candidat.pk).exists())
        self.assertFalse(candidat.is_active)
        self.assertFalse(candidat.profil.en_attente)

    def test_interface_admin_ne_permet_pas_la_suppression_des_comptes(self):
        requete = RequestFactory().get("/admin/auth/user/")
        requete.user = self.admin

        self.assertFalse(admin.site._registry[User].has_delete_permission(requete))

    def test_compte_administe_rejette_email_utilise_comme_identifiant(self):
        User.objects.create_user("occupe@example.test", "autre@example.test", "MotDePasse!2026")
        self.client.force_login(self.admin)

        reponse = self.client.post(
            reverse("inviter_utilisateur"),
            {
                "first_name": "Nouveau",
                "last_name": "Compte",
                "email": "occupe@example.test",
                "role": Profil.Role.AGENT,
            },
        )

        self.assertEqual(reponse.status_code, 302)
        self.assertFalse(User.objects.filter(email="occupe@example.test").exists())

    def test_reprise_email_notification_anomalie_et_respect_du_retrait(self):
        assigne = User.objects.create_user("assigne", "assigne@example.test", "MotDePasse!2026")
        assigne.profil.role = Profil.Role.AGENT
        assigne.profil.secteur = self.nord
        assigne.profil.recevoir_notifications = False
        assigne.profil.save()
        anomalie = Anomalie.objects.create(
            categorie=self.taxes,
            secteur=self.nord,
            mois=date(2026, 1, 1),
            description="Notification à reprendre",
            assigne_a=assigne,
        )

        with patch("recettes.notifications.send_mail", side_effect=smtplib.SMTPException("SMTP indisponible")):
            with self.assertLogs("recettes.notifications", level="ERROR"):
                notifier_anomalies_en_attente()
        notification_assigne = Notification.objects.get(
            destinataire=assigne, type=Notification.Type.ANOMALIE
        )
        self.assertFalse(notification_assigne.envoyee_par_email)

        with patch("recettes.notifications.send_mail", return_value=1):
            reprises = notifier_anomalies_en_attente()

        self.assertGreater(reprises, 0)
        self.assertTrue(
            Notification.objects.filter(
                type=Notification.Type.ANOMALIE, envoyee_par_email=True
            ).exists()
        )

    def test_rappels_respectent_la_frequence_de_chaque_secteur(self):
        self.nord.frequence_versement = Secteur.Frequence.SEPT
        self.nord.save(update_fields=["frequence_versement"])
        self.sud.frequence_versement = Secteur.Frequence.TRENTE
        self.sud.save(update_fields=["frequence_versement"])

        notifier_echeances_recouvrement(date(2026, 10, 10))

        self.assertTrue(Notification.objects.filter(cle_unicite__contains=f"2026-09-{self.nord.pk}").exists())
        self.assertFalse(Notification.objects.filter(cle_unicite__contains=f"2026-09-{self.sud.pk}").exists())

        notifier_echeances_recouvrement(date(2026, 10, 30))

        self.assertTrue(Notification.objects.filter(cle_unicite__contains=f"2026-09-{self.sud.pk}").exists())

    def test_email_de_relance_reussi_est_comptabilise(self):
        with patch("recettes.notifications.send_mail", side_effect=smtplib.SMTPException("SMTP indisponible")):
            with self.assertLogs("recettes.notifications", level="ERROR"):
                self.assertTrue(
                    _creer_et_envoyer(
                        self.admin, Notification.Type.SYSTEME, "Test", "Message", "REPRISE-EMAIL"
                    )
                )

        with patch("recettes.notifications.send_mail", return_value=1):
            self.assertTrue(
                _creer_et_envoyer(
                    self.admin, Notification.Type.SYSTEME, "Test", "Message", "REPRISE-EMAIL"
                )
            )

    def test_distribution_student_et_degres_de_liberte(self):
        self.assertEqual(_t_critique_95(1), 12.7062)
        self.assertEqual(_t_critique_95(10), 2.2281)
        with self.assertRaises(ValueError):
            _t_critique_95(0)

        from .prevision import _modele

        historique = [date(2023 + i // 12, i % 12 + 1, 1) for i in range(36)]
        montants = [1000 + 5 * i + (100 if mois.month == 12 else 0) for i, mois in enumerate(historique)]
        _, prevus, bas, hauts, saisonnalite = _modele(historique, montants, 2)

        self.assertTrue(saisonnalite)
        self.assertTrue(all(basse <= prevu <= haute for basse, prevu, haute in zip(bas, prevus, hauts)))

    def test_un_fichier_de_rapport_absent_renvoie_404(self):
        rapport = RapportGenere.objects.create(
            type_doc=RapportGenere.TypeDoc.SYNTHESE,
            format=RapportGenere.Format.PDF,
            periode="Exercice 2026",
            cree_par=self.responsable,
        )
        rapport.fichier.name = "rapports/2026/01/introuvable.pdf"
        rapport.save(update_fields=["fichier"])
        self.client.force_login(self.responsable)

        reponse = self.client.get(reverse("telecharger_rapport", args=[rapport.pk]))

        self.assertEqual(reponse.status_code, 404)

    def test_rapports_pdf_et_excel_sont_generes(self):
        self.client.force_login(self.responsable)
        pdf = self.client.post(
            reverse("rapports"), {"doc_type": "SYNTHESE", "format": "PDF"}
        )
        excel = self.client.post(
            reverse("rapports"), {"doc_type": "SYNTHESE", "format": "XLSX"}
        )

        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(excel.status_code, 200)
        self.assertTrue(excel.content.startswith(b"PK"))

    def test_les_valeurs_xlsx_commencant_par_une_formule_sont_neutralisees(self):
        from openpyxl import load_workbook

        mois = timezone.localdate().replace(day=1)
        Recette.objects.create(
            categorie=self.taxes,
            secteur=self.nord,
            mois=mois,
            date_imputation=timezone.localdate(),
            montant=Decimal("100"),
            libelle=" =1+1",
        )

        classeur = load_workbook(generer_rapport_xlsx(annee=mois.year, mois=mois))
        valeurs = [
            cellule
            for feuille in classeur.worksheets
            for ligne in feuille.iter_rows()
            for cellule in ligne
            if cellule.value == "' =1+1"
        ]

        self.assertEqual(len(valeurs), 1)
        self.assertNotEqual(valeurs[0].data_type, "f")

    def test_texte_sur_neutralise_les_formules_precedees_despaces(self):
        self.assertEqual(texte_sur("  =1+1"), "'  =1+1")
        self.assertEqual(texte_sur(" -123"), " -123")

    def test_commande_notifications_accepte_la_reprise_des_anomalies(self):
        call_command("envoyer_notifications", type="anomalies", verbosity=0)

    def test_commande_creer_admin_cree_et_synchronise_le_role(self):
        with patch.dict(
            "os.environ",
            {
                "DJANGO_SUPERUSER_USERNAME": "admin-render-test",
                "DJANGO_SUPERUSER_EMAIL": "admin-render@example.test",
                "DJANGO_SUPERUSER_PASSWORD": "MotDePasse!2026",
            },
        ):
            call_command("creer_admin", verbosity=0)
            call_command("creer_admin", verbosity=0)

        utilisateur = User.objects.get(username="admin-render-test")
        self.assertTrue(utilisateur.is_superuser)
        self.assertTrue(utilisateur.is_staff)
        self.assertEqual(utilisateur.profil.role, Profil.Role.ADMIN)

    @patch("recettes.management.commands.boucle_taches.historiser_previsions", return_value=1)
    @patch("recettes.management.commands.boucle_taches.notifier_anomalies_en_attente", return_value=2)
    @patch("recettes.management.commands.boucle_taches.notifier_echeances_recouvrement", return_value=3)
    @patch("recettes.management.commands.boucle_taches.notifier_ecarts_significatifs", return_value=4)
    def test_commande_planifiee_execute_une_iteration_complete(
        self, ecarts, echeances, anomalies, previsions
    ):
        call_command("boucle_taches", une_fois=True, verbosity=0)

        ecarts.assert_called_once()
        echeances.assert_called_once()
        anomalies.assert_called_once()
        previsions.assert_called_once()
